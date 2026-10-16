"""QR-Code für die Lernkarten KB I erzeugen und (optional) ans Ende der Ich-kann-Liste setzen.

Aufruf (im Hauptverzeichnis des Repositorys):
    pip install segno python-docx zxing-cpp pillow
    python werkzeuge/qr_kb1_lernkarten.py                      # nur qr/qr_kb1_lernkarten.png
    python werkzeuge/qr_kb1_lernkarten.py Ich_kann_Liste.docx  # zusätzlich Ich_kann_Liste_mit_QR.docx

- Liest Besitzer und Namen des Repositorys aus dem Git-Remote „origin“;
  die Adresse ist https://<besitzer>.github.io/<repository>/kb1/
- Erzeugt qr/qr_kb1_lernkarten.png (Fehlerkorrektur M, Ruhezone 4 Module,
  mindestens 600 × 600 px) und prüft, dass er sich dekodieren lässt.
- Mit einer .docx-Datei: hängt unten einen kleinen Kasten mit QR-Code (3,2 cm) und
  kurzem Text an und speichert das Ergebnis als „…_mit_QR.docx“. Das Original bleibt unverändert.
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from qr_einsetzen import qr_erzeugen, qr_pruefen, repo_aus_remote  # noqa: E402

WURZEL = Path(__file__).resolve().parent.parent
QR_DATEI = WURZEL / "qr" / "qr_kb1_lernkarten.png"


def zelle_rahmen(zelle, farbe="3F6C72"):
    tc_pr = zelle._tc.get_or_add_tcPr()
    raender = OxmlElement("w:tcBorders")
    for seite in ("top", "left", "bottom", "right"):
        r = OxmlElement(f"w:{seite}")
        r.set(qn("w:val"), "single")
        r.set(qn("w:sz"), "12")
        r.set(qn("w:color"), farbe)
        raender.append(r)
    tc_pr.append(raender)


def in_liste_einsetzen(docx_pfad, adresse):
    doc = Document(docx_pfad)
    doc.add_paragraph()
    tabelle = doc.add_table(rows=1, cols=2)
    tabelle.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabelle.autofit = False
    links, rechts = tabelle.rows[0].cells
    links.width, rechts.width = Cm(3.8), Cm(11.5)
    for z in (links, rechts):
        zelle_rahmen(z)

    p = links.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(QR_DATEI), width=Cm(3.2))

    p = rechts.paragraphs[0]
    r = p.add_run("Lernkarten fürs Smartphone")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x25, 0x43, 0x4B)
    p = rechts.add_paragraph(
        "QR-Code scannen und alle Kompetenzen dieser Liste als Karteikarten lernen: "
        "vorne die Kompetenz, hinten die Beispielantwort. Ohne Anmeldung – "
        "Ihr Lernstand bleibt auf Ihrem Handy."
    )
    p = rechts.add_paragraph()
    r = p.add_run(adresse.replace("https://", "").rstrip("/"))
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x3F, 0x6C, 0x72)
    for zelle in (links, rechts):
        zelle.vertical_alignment = 1  # Mitte
    doc.add_paragraph()

    ziel = Path(docx_pfad).with_name(Path(docx_pfad).stem + "_mit_QR.docx")
    doc.save(ziel)
    return ziel


def main():
    besitzer, repo = repo_aus_remote()
    adresse = f"https://{besitzer.lower()}.github.io/{repo}/kb1/"
    QR_DATEI.parent.mkdir(exist_ok=True)
    print(f"Adresse: {adresse}")
    qr_erzeugen(adresse, QR_DATEI)
    qr_pruefen(QR_DATEI, adresse)
    print(f"  gespeichert: {QR_DATEI.relative_to(WURZEL)}")
    if len(sys.argv) > 1:
        ziel = in_liste_einsetzen(sys.argv[1], adresse)
        print(f"  Ich-kann-Liste mit QR-Code: {ziel}")


if __name__ == "__main__":
    main()
